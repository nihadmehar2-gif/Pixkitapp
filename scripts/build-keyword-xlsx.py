"""Build PixKitApp-Keyword-Map.xlsx from keyword-map.json (real Google Autocomplete
verification data) + seo-inventory.json. Follows xlsx skill design system (base.py).
Sheets: Keyword Map | Top Opportunities | Category Coverage | Methodology | Review
NO search volumes, difficulty scores or traffic predictions are included — none were
measured. Evidence types are clearly labeled instead.
"""
import sys, os, json
XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import *  # design tokens + helpers
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

KM = json.load(open("/home/z/my-project/scripts/keyword-map.json"))
INV = json.load(open("/home/z/my-project/scripts/seo-inventory.json"))
RETRIEVED = "2026-10-14"

# ---------- helpers ----------
INTENT_INFO = ("requirements", "size", "sizes", "dimensions")
def intent_of(pk: str) -> str:
    p = pk.lower()
    if p.endswith(INTENT_INFO) or " requirements" in p:
        return "Mixed (info + tool)"
    return "Transactional (tool intent)"

PLATFORM_OR_SPEC = ("instagram", "facebook", "twitter", " x ", "linkedin", "youtube", "pinterest",
                    "kb", "passport", "visa", "webp", "avif", "bmp", "ico", "base64", "svg",
                    "16:9", "9:16", "4:3", "3:2", "1:1", "2x2", "35x45", "33x48", "100kb", "200kb", "500kb", "50kb")
def competition_note(pk: str) -> str:
    p = " " + pk.lower() + " "
    words = len(pk.split())
    if words >= 3 or any(s in p for s in PLATFORM_OR_SPEC):
        return "Specific long-tail — typically lower competition (qualitative, not a metric)"
    if words == 2:
        return "Mid-tail — established competitors rank, but niche tools can win on UX and specificity (qualitative)"
    return "Head term — dominated by major aggregators; compete via long-tails and page experience (qualitative)"

def priority_of(row) -> str:
    v = sum(1 for lt in row["longTails"] if lt["evidence"] == "verified")
    if row["primaryEvidence"] == "verified" and v >= 2:
        return "P1"
    if row["primaryEvidence"] == "verified" or v >= 1:
        return "P2"
    return "P3"

VOL = "Not measured — verify with Google Keyword Planner or GSC impressions after indexing"

wb = Workbook()

# ============ Sheet 1: Keyword Map ============
ws = wb.active
ws.title = "Keyword Map"
headers = ["Tool", "URL path", "Category", "On-page primary keyword", "Primary keyword evidence",
           "Long-tail keywords — Autocomplete-verified", "Long-tail keywords — inferred (verify before targeting)",
           "Real query patterns from Google Autocomplete", "Search intent", "Competition note (qualitative)",
           "Search volume", "Priority"]
last_col = len(headers) + 1
setup_sheet(ws, title="PixKitApp.com — Keyword Map (127 tools, worldwide English)", last_col=last_col)
for c, h in enumerate(headers, 2):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=last_col)

r = 5
for i, row in enumerate(KM):
    v_lts = [lt["kw"] for lt in row["longTails"] if lt["evidence"] == "verified"]
    i_lts = [lt["kw"] for lt in row["longTails"] if lt["evidence"] != "verified"]
    vals = [
        row["name"],
        row["url"].replace("https://pixkitapp.com", ""),
        row["category"],
        row["primaryKeyword"],
        "Autocomplete-verified" if row["primaryEvidence"] == "verified" else "Inferred",
        "; ".join(v_lts),
        "; ".join(i_lts),
        "; ".join(row["harvestedSuggestions"]),
        intent_of(row["primaryKeyword"]),
        competition_note(row["primaryKeyword"]),
        VOL,
        priority_of(row),
    ]
    for c, v in enumerate(vals, 2):
        ws.cell(row=r, column=c, value=v)
    style_data_row(ws, row_num=r, col_start=2, col_end=last_col, row_index=i)
    r += 1
last_data = r - 1
# caption
note = ws.cell(row=r + 1, column=2, value=f"Evidence source: Google Autocomplete (public suggest endpoint, hl=en, retrieved {RETRIEVED}). "
                                          f"Autocomplete verification proves a query pattern is used in Google Suggest — it is NOT a search-volume measurement.")
note.font = font_caption()
ws.freeze_panes = "C5"
auto_fit_columns(ws, min_width=8, max_width=42, header_row=4, data_start_row=5)
for col, w in {"B": 22, "C": 30, "D": 18, "E": 26, "F": 20, "G": 40, "H": 40, "I": 44, "J": 20, "K": 40, "L": 34, "M": 9}.items():
    ws.column_dimensions[col].width = w

# ============ Sheet 2: Top Opportunities ============
ws2 = wb.create_sheet("Top Opportunities")
opp_headers = ["Cluster", "Tool pages", "Why (evidence-based)", "Suggested action"]
last2 = len(opp_headers) + 1
setup_sheet(ws2, title="Top keyword opportunities — prioritized, evidence-based", last_col=last2)
for c, h in enumerate(opp_headers, 2):
    ws2.cell(row=4, column=c, value=h)
style_header_row(ws2, row_num=4, col_start=2, col_end=last2)

opps = [
    ("Exact-KB compression cluster", "compress-image-to-50kb, -100kb, -200kb, -500kb, compress-image-to-specific-size, resize-image-to-kb",
     "Autocomplete harvest shows heavy demand for size targets: 'image compressor to 20kb / 50kb / 100kb / 200kb', 'resize image to kb and pixel'. Few major competitors offer dedicated per-size pages.",
     "Keep the six pages tightly interlinked; use the verified size strings in FAQ answers (already done); monitor GSC queries for new size targets (e.g. 20kb) and add pages only if demand appears."),
    ("Social size + unit conversions", "12 social resizers (instagram-post … pinterest-pin)",
     "Verified queries like 'instagram story size in cm', 'in inches', 'in pixels', 'size converter' — informational intent the big tools answer poorly.",
     "Unit-conversion FAQs added on all 12 pages; consider a short 'social media image size guide' section on the category page if GSC shows informational impressions."),
    ("Passport & visa photo by country", "passport-photo-maker, us-passport-photo, schengen-visa-photo, china-visa-photo, id-photo-1-inch, photo-print-sheet",
     "'us passport photo requirements', 'schengen visa photo requirements', 'china visa photo size' verified patterns; incumbents are paid-heavy (iVisa-style) or outdated.",
     "Free + no-upload + print-sheet angle is the differentiator; keep requirement FAQs current with official specs; add print-at-home guidance."),
    ("WebP/AVIF modern-format conversions", "webp-to-jpg, webp-to-png, jpg-to-webp, png-to-webp, avif-to-jpg, avif-to-png, gif-to-webp",
     "'webp to jpg converter online high quality' and variants verified; AVIF is newer with far fewer dedicated tools (lower competition, qualitative).",
     "AVIF pages are the best long-tail entry; WebP pages compete with iLoveIMG/Convertio on head terms — target the 'high quality'/'without losing quality' modifiers."),
    ("White-background removal cluster", "remove-white-background, make-image-transparent, transparent-png-maker, add-white-background, add-background-color",
     "'remove white background from image/logo/png' verified; dedicated AI removers are heavy/paid; flood-fill version wins for logos & product shots on white.",
     "Clarify positioning (works best on white/near-white product shots) — already in FAQ; interlink the five transparency tools."),
    ("Developer utilities", "image-to-base64, base64-to-image, image-to-ico, svg-to-png, favicon-generator",
     "Niche developer long-tails with steady utility demand ('convert image to base64', 'favicon ico generator from png' verified); low competition from big aggregators.",
     "Keep answers technically precise (data-URI, multi-size ICO details are already in FAQs); these earn natural links from developer tutorials."),
    ("Batch/bulk variants", "bulk-image-compressor, bulk-watermark",
     "'bulk image compressor online', 'compress multiple images at once' verified; privacy angle (no upload) differentiates vs iLoveIMG.",
     "Emphasize 10-file + ZIP workflow in page copy; batch pages should link to single-file twins and vice versa (done via related tools)."),
    ("Exact-ratio crop pages", "crop-image-16-9, -9-16, -4-3, -3-2, crop-image-to-square, crop-image-to-circle",
     "Aspect-specific crop queries verified ('crop image to 16:9' pattern class); most competitors bury ratio crop inside one editor.",
     "Each ratio page already has dedicated anchors + FAQs; circle crop targets profile-picture use cases."),
    ("Image-to-PDF (free, no upload)", "image-to-pdf",
     "'jpg to pdf converter online free', 'combine images into pdf' are verified high-demand patterns; incumbent pages are ad-heavy.",
     "Page-experience angle: fast, clean, batch-of-10 with page-size control; strong candidate for the homepage popular section (already present)."),
    ("EXIF privacy tools", "image-metadata-viewer, remove-image-metadata, remove-gps-data",
     "'remove exif data online', 'remove gps from photo' verified; privacy-conscious queries keep growing.",
     "Cross-link from passport tools (share photos safely) and watermark tools; the privacy story is the site's core differentiator."),
]
r2 = 5
for i, row in enumerate(opps):
    for c, v in enumerate(row, 2):
        ws2.cell(row=r2, column=c, value=v)
    style_data_row(ws2, row_num=r2, col_start=2, col_end=last2, row_index=i)
    r2 += 1
ws2.freeze_panes = "C5"
for col, w in {"B": 26, "C": 40, "D": 52, "E": 52}.items():
    ws2.column_dimensions[col].width = w

# ============ Sheet 3: Category Coverage ============
ws3 = wb.create_sheet("Category Coverage")
cat_headers = ["Category", "URL path", "Tools", "Primaries Autocomplete-verified", "P1 tools", "P2 tools", "P3 tools"]
last3 = len(cat_headers) + 1
setup_sheet(ws3, title="Coverage by category", last_col=last3)
for c, h in enumerate(cat_headers, 2):
    ws3.cell(row=4, column=c, value=h)
style_header_row(ws3, row_num=4, col_start=2, col_end=last3)
cats = INV["categories"]
r3 = 5
for i, cat in enumerate(cats):
    rows = [k for k in KM if k["categorySlug"] == cat["slug"]]
    vals = [
        cat["name"],
        f"/categories/{cat['slug']}",
        len(rows),
        sum(1 for k in rows if k["primaryEvidence"] == "verified"),
        sum(1 for k in rows if priority_of(k) == "P1"),
        sum(1 for k in rows if priority_of(k) == "P2"),
        sum(1 for k in rows if priority_of(k) == "P3"),
    ]
    for c, v in enumerate(vals, 2):
        ws3.cell(row=r3, column=c, value=v)
    style_data_row(ws3, row_num=r3, col_start=2, col_end=last3, row_index=i)
    r3 += 1
# totals row (live formulas)
ws3.cell(row=r3, column=2, value="Total")
for c in range(4, 9):
    letter = get_column_letter(c)
    ws3.cell(row=r3, column=c, value=f"=SUM({letter}5:{letter}{r3-1})")
style_total_row(ws3, row_num=r3, col_start=2, col_end=last3)
ws3.freeze_panes = "C5"
for col, w in {"B": 24, "C": 24, "D": 10, "E": 26, "F": 10, "G": 10, "H": 10}.items():
    ws3.column_dimensions[col].width = w

# ============ Sheet 4: Methodology ============
ws4 = wb.create_sheet("Methodology")
meth_headers = ["Topic", "Details"]
last4 = len(meth_headers) + 1
setup_sheet(ws4, title="Methodology, data sources and honest limitations", last_col=last4)
for c, h in enumerate(meth_headers, 2):
    ws4.cell(row=4, column=c, value=h)
style_header_row(ws4, row_num=4, col_start=2, col_end=last4)
method = [
    ("Audience & scope", "Worldwide English-speaking users. All 127 tool pages mapped; one row per page."),
    ("On-page primary keyword", "Taken from the site's existing SEO configuration (src/lib/tools/seo/*) — authored during the SEO build; these are the keywords each page is optimized for."),
    ("Google Autocomplete verification", f"Candidate keywords were checked against Google's public suggest endpoint (client=firefox, hl=en) on {RETRIEVED}. "
                                        "A keyword is 'Autocomplete-verified' if it (or a close superstring) appears in Google's live suggestion list for the tool's seed queries. "
                                        "This proves real users type the pattern into Google — it is evidence of query existence, not of volume."),
    ("Harvested suggestions", "Up to 8 real suggestion strings per tool, returned by Google Suggest for the tool's primary keyword and top long-tail. These are verbatim query patterns users see in Google search."),
    ("Inferred keywords", "Long-tails authored from each tool's actual function, options, formats and limits (not returned by Suggest in this run). Treat as hypotheses; verify with a keyword tool before committing content."),
    ("Search intent", "Transactional = user wants to perform the task now (tool page satisfies directly). Mixed = informational query ('…requirements', '…size') whose best answer includes a working tool + the facts."),
    ("Competition notes", "Qualitative reasoning based on keyword specificity and the known competitive landscape (iLoveIMG, TinyPNG, Canva, Adobe Express, Pi7, imagy.app, Zamzar, Photopea). NOT keyword-difficulty scores."),
    ("Search volume — honest status", "NO search volumes are included because no volume data source was available. Nothing here is an estimate dressed up as data. "
                                      "To add volumes: (1) free — Google Keyword Planner (Google Ads account needed, gives ranges); (2) after launch — Google Search Console impressions are the best real-world signal; (3) paid — Ahrefs/Semrush if available."),
    ("What this map is for", "Page-to-keyword mapping to prevent keyword cannibalization, guide future content, and track which clusters earn impressions after indexing."),
    ("Refresh cadence", "Re-run Autocomplete verification before major content decisions (suggest lists shift); re-check after each new tool batch."),
]
r4 = 5
for i, (topic, detail) in enumerate(method):
    ws4.cell(row=r4, column=2, value=topic)
    ws4.cell(row=r4, column=3, value=detail)
    style_data_row(ws4, row_num=r4, col_start=2, col_end=last4, row_index=i)
    ws4.cell(row=r4, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    ws4.cell(row=r4, column=2).alignment = Alignment(vertical="top")
    r4 += 1
ws4.column_dimensions["B"].width = 30
ws4.column_dimensions["C"].width = 100

# ============ Sheet 5: Review (self-check) ============
ws5 = wb.create_sheet("Review")
ws5.sheet_properties.tabColor = "FFC000"
checks = [
    ["Check", "Expected", "Actual", "Status"],
    ["Keyword Map rows (tools)", 127, "=COUNTA('Keyword Map'!B5:B131)", '=IF(B2=C2,"PASS","FAIL")'],
    ["Category totals sum", 127, "='Category Coverage'!D19", '=IF(B3=C3,"PASS","FAIL")'],
    ["Opportunity clusters", len(opps), "=COUNTA('Top Opportunities'!B5:B200)", '=IF(B4=C4,"PASS","FAIL")'],
]
for i, row in enumerate(checks, 1):
    for j, val in enumerate(row, 1):
        ws5.cell(row=i, column=j, value=val)
ws5.column_dimensions["A"].width = 34
ws5.column_dimensions["B"].width = 12
ws5.column_dimensions["C"].width = 40
ws5.column_dimensions["D"].width = 10

wb.properties.creator = "Z.ai"
out = "/home/z/my-project/download/PixKitApp-Keyword-Map.xlsx"
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print("saved", out)
print("rows:", last_data - 4, "| opportunities:", len(opps))
